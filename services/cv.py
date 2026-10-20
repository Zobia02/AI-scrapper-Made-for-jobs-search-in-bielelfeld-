from io import BytesIO
from zipfile import ZipFile, BadZipFile

from docx import Document
from pypdf import PdfReader

from matching.extraction import skills_in


def extract_cv(data: bytes, filename: str) -> str:
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("CV must be 8 MB or smaller.")
    try:
        if filename.lower().endswith(".pdf"):
            if not data.startswith(b"%PDF"):
                raise ValueError("This is not a valid PDF.")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Upload an unencrypted PDF.")
            if len(reader.pages) > 30:
                raise ValueError("CV is limited to 30 pages.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif filename.lower().endswith(".docx"):
            with ZipFile(BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 30_000_000:
                    raise ValueError("Expanded DOCX is too large.")
            doc = Document(BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        else:
            raise ValueError("Use PDF or DOCX.")
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("The CV could not be read. Try exporting it again.") from exc
    if not text.strip():
        raise ValueError("No text found. Scanned PDFs need OCR before upload; paste corrected text manually.")
    return text[:100_000]


def parsed_skills(text):
    return skills_in(text)
