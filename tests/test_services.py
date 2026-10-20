from io import BytesIO

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

from matching.profile import Profile
from matching.scorer import score_job
from services.cv import extract_cv
from services.exports import csv_bytes, excel_bytes, jobs_frame
from services.filters import filter_jobs
from services.llm import request_advice


def test_csv_excel_and_injection(job):
    data = job.model_dump(mode="json")
    data["title"] = "=HYPERLINK(\"bad\")"
    data["match"] = score_job(data, Profile())
    frame = jobs_frame([data])
    parsed = pd.read_csv(BytesIO(csv_bytes(frame)))
    assert parsed.iloc[0]["title"].startswith("'=")
    assert parsed.iloc[0]["apply_url"] == str(job.apply_url)
    workbook = load_workbook(BytesIO(excel_bytes(frame)))
    cell = workbook.active["B2"]
    assert cell.data_type != "f"


def test_filters_do_not_invent_dates_or_pay(job):
    data = job.model_dump(mode="json")
    data["match"] = score_job(data, Profile())
    assert filter_jobs([data], {"days": 7}) == []
    data["paid"] = "Unknown"
    assert filter_jobs([data], {"paid_only": True}) == []
    assert filter_jobs([data], {"languages": []}) == []
    assert len(filter_jobs([data], {})) == 1


def test_cv_docx_local():
    document = Document()
    document.add_paragraph("TEST CV ONLY: Python and RAG developer.")
    stream = BytesIO()
    document.save(stream)
    assert "Python" in extract_cv(stream.getvalue(), "cv.docx")


def test_cv_pdf_local():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 40 740 Td (TEST CV ONLY: Python and RAG developer.) Tj ET")
    page[NameObject("/Contents")] = content
    stream = BytesIO()
    writer.write(stream)
    assert "Python and RAG" in extract_cv(stream.getvalue(), "test-cv.pdf")


def test_cv_invalid_and_oversize():
    with pytest.raises(ValueError):
        extract_cv(b"not a pdf", "cv.pdf")
    with pytest.raises(ValueError):
        extract_cv(b"x" * (9 * 1024 * 1024), "cv.pdf")


def test_llm_requires_consent():
    with pytest.raises(ValueError, match="authorization"):
        request_advice({}, consent=False)
